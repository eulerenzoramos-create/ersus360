"""
PMS Vivo — regras de negócio: perfis, planilha modelo (Excel) com revisão humana, versões do PMS,
situação das ações e validação da execução.

Regras (determinação do usuário):
  • Nada é inventado: célula vazia ou ilegível na planilha fica vazia (com aviso/erro), nunca estimada.
  • Versão aprovada do PMS não é editada: alterar = nova versão (proposta), que só vale depois de aprovada;
    a aprovada anterior passa a "substituido" e continua consultável.
  • A situação OFICIAL de uma ação só muda quando a atualização informada é validada por outra pessoa
    com perfil de validação ("não basta marcar uma caixa"). Cancelamento exige justificativa e autorização.
"""
from __future__ import annotations

import io
import re
import uuid
from datetime import date, datetime

from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from models.pms_vivo import (PmsAcao, PmsAcaoMeta, PmsAnexo, PmsAtualizacao, PmsDiretriz, PmsHistorico,
                             PmsMeta, PmsMetaAnual, PmsObjetivo, PmsPas, PmsPlano)

# ── Perfis (reaproveita os perfis existentes do ERSUS360; ver models/usuario.Perfil) ──────────────
#   Secretário Municipal de Saúde → gestor · Coordenação de Planejamento → planejamento
#   Coordenadores de áreas → coordenador · Assessoria técnica → admin/auditoria
#   Consulta e controle social → consulta/conselho (somente leitura)
#   Responsável pela execução → o usuário indicado na ação (qualquer perfil do município)
PERFIS_PLANEJAMENTO = {"admin", "gestor", "planejamento"}               # cadastram/importam/aprovam
PERFIS_VALIDACAO = {"admin", "gestor", "planejamento", "coordenador"}   # validam a execução
PERFIS_AUTORIZA_CANCELAMENTO = {"admin", "gestor", "planejamento"}

TIPOS_INDICADOR = {
    "percentual": "Percentual (não somar entre anos)",
    "taxa": "Taxa / coeficiente (não somar entre anos)",
    "indice": "Índice / razão (não somar entre anos)",
    "absoluto_acumulavel": "Número absoluto acumulável (soma entre anos)",
    "absoluto_nao_acumulavel": "Número absoluto não acumulável (estoque, ex.: equipes existentes)",
    "binario": "Sim/Não (ex.: plano elaborado)",
}
POLARIDADES = {"maior_melhor": "Quanto maior, melhor", "menor_melhor": "Quanto menor, melhor"}
SITUACOES_ACAO = {
    "nao_iniciada": "Não iniciada", "em_andamento": "Em andamento", "concluida": "Concluída",
    "parcial": "Parcialmente executada", "nao_executada": "Não executada", "cancelada": "Cancelada",
}
EXIGE_JUSTIFICATIVA = {"parcial", "nao_executada", "cancelada"}
SITUACOES_ENCERRADAS = {"concluida", "parcial", "nao_executada", "cancelada"}
PRAZO_PROXIMO_DIAS = 15


def pode_planejar(u) -> bool:
    return bool(u.administrador_geral or u.role in PERFIS_PLANEJAMENTO)


def pode_validar(u) -> bool:
    return bool(u.administrador_geral or u.role in PERFIS_VALIDACAO)


def historico(db: AsyncSession, municipio_id: int, entidade: str, entidade_id: int, evento: str,
              descricao: str, usuario: str | None) -> None:
    db.add(PmsHistorico(municipio_id=municipio_id, entidade=entidade, entidade_id=entidade_id, evento=evento,
                        descricao=descricao[:4000], usuario=usuario))


# ── Planilha modelo e leitura com revisão ────────────────────────────────────

COLUNAS_FIXAS = [
    ("diretriz_codigo", "Diretriz (código)"), ("diretriz", "Diretriz"),
    ("objetivo_codigo", "Objetivo (código)"), ("objetivo", "Objetivo"),
    ("meta_codigo", "Meta (código)"), ("meta", "Meta"),
    ("indicador", "Indicador"), ("tipo_indicador", "Tipo do indicador"), ("unidade", "Unidade de medida"),
    ("polaridade", "Polaridade"), ("formula", "Fórmula de cálculo"),
    ("linha_base", "Linha de base"), ("ano_linha_base", "Ano da linha de base"),
    ("meta_quadrienal", "Meta quadrienal"), ("fonte", "Fonte de informação"),
    ("setor_responsavel", "Setor responsável"),
]
OBRIGATORIAS = ("diretriz_codigo", "diretriz", "objetivo_codigo", "objetivo", "meta_codigo", "meta")


def _norm(s) -> str:
    import unicodedata
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def modelo_xlsx(anos: list[int]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "PMS"
    cab = [r for _, r in COLUNAS_FIXAS] + [f"Meta {a}" for a in anos]
    ws.append(cab)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DBEAFE")
    for i, _ in enumerate(cab, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = 22
    tipos = DataValidation(type="list", formula1='"' + ",".join(TIPOS_INDICADOR) + '"', allow_blank=True)
    polar = DataValidation(type="list", formula1='"' + ",".join(POLARIDADES) + '"', allow_blank=True)
    ws.add_data_validation(tipos)
    ws.add_data_validation(polar)
    tipos.add("H2:H2000")
    polar.add("J2:J2000")
    inst = wb.create_sheet("Instruções")
    for linha in [
        ["Uma linha por meta. Repita o código e o texto da diretriz e do objetivo em cada linha."],
        ["Preencha só o que está no PMS aprovado. Deixe em branco o que o documento não traz — nada é estimado."],
        ["Números: use vírgula ou ponto decimal; '%' é aceito e ignorado (ex.: 85% → 85)."],
        ["Tipos do indicador: " + "; ".join(f"{k} = {v}" for k, v in TIPOS_INDICADOR.items())],
        ["Polaridade: maior_melhor ou menor_melhor."],
        ["Depois de enviar, o ERSUS360 mostra a prévia para revisão; nada é gravado antes da confirmação."],
    ]:
        inst.append(linha)
    inst.column_dimensions["A"].width = 140
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def numero(v) -> tuple[float | None, bool]:
    """(valor, ok). Vazio → (None, True). Texto que não é número → (None, False) — nunca estima."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None, True
    if isinstance(v, bool):
        return None, False
    if isinstance(v, (int, float)):
        return float(v), True
    s = str(v).strip().replace("%", "").replace(" ", "")
    if re.fullmatch(r"-?\d{1,3}(\.\d{3})+(,\d+)?", s):      # 1.234,5
        s = s.replace(".", "").replace(",", ".")
    else:
        s = s.replace(",", ".")
    try:
        return float(s), True
    except ValueError:
        return None, False


def ler_planilha(conteudo: bytes, anos: list[int]) -> dict:
    """Lê a planilha do PMS e devolve a estrutura para REVISÃO (não grava nada)."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except Exception:
        raise ValueError("Arquivo não é uma planilha Excel (.xlsx) válida.")
    ws = wb["PMS"] if "PMS" in wb.sheetnames else wb.worksheets[0]
    linhas = list(ws.iter_rows(values_only=True))
    if not linhas:
        raise ValueError("Planilha vazia.")
    cab = [_norm(c) for c in linhas[0]]
    mapa: dict[str, int] = {}
    for chave, rotulo in COLUNAS_FIXAS:
        if _norm(rotulo) in cab:
            mapa[chave] = cab.index(_norm(rotulo))
    for a in anos:
        if _norm(f"Meta {a}") in cab:
            mapa[f"ano_{a}"] = cab.index(_norm(f"Meta {a}"))
    faltam = [r for k, r in COLUNAS_FIXAS if k in OBRIGATORIAS and k not in mapa]
    if faltam:
        raise ValueError("Colunas obrigatórias ausentes: " + ", ".join(faltam) + ". Use a planilha modelo.")

    erros, avisos = [], []
    diretrizes: dict[str, dict] = {}
    codigos_meta: set[str] = set()
    for n, row in enumerate(linhas[1:], start=2):
        val = {k: (row[i] if i < len(row) else None) for k, i in mapa.items()}
        if all(v in (None, "") for v in val.values()):
            continue
        txt = {k: (str(v).strip() if v not in (None, "") else None) for k, v in val.items()}
        falta = [r for k, r in COLUNAS_FIXAS if k in OBRIGATORIAS and not txt.get(k)]
        if falta:
            erros.append(f"Linha {n}: falta {', '.join(falta)}.")
            continue
        meta = {"linha": n, "codigo": txt["meta_codigo"], "descricao": txt["meta"]}
        for campo in ("indicador", "unidade", "formula", "fonte", "setor_responsavel"):
            meta[campo] = txt.get(campo)
        tipo = (txt.get("tipo_indicador") or "").lower() or None
        if tipo and tipo not in TIPOS_INDICADOR:
            erros.append(f"Linha {n}: tipo do indicador '{tipo}' inválido.")
            tipo = None
        if not tipo:
            avisos.append(f"Linha {n} (meta {meta['codigo']}): tipo do indicador não informado — "
                          "o acumulado anual não será calculado até ser definido.")
        meta["tipo_indicador"] = tipo
        pol = (txt.get("polaridade") or "").lower() or None
        if pol and pol not in POLARIDADES:
            erros.append(f"Linha {n}: polaridade '{pol}' inválida.")
            pol = None
        meta["polaridade"] = pol
        for campo in ("linha_base", "meta_quadrienal"):
            v, ok = numero(val.get(campo))
            if not ok:
                erros.append(f"Linha {n}: '{val.get(campo)}' em {dict(COLUNAS_FIXAS)[campo]} não é número.")
            meta[campo] = v
        v, ok = numero(val.get("ano_linha_base"))
        meta["ano_linha_base"] = int(v) if ok and v is not None else None
        if not ok:
            erros.append(f"Linha {n}: ano da linha de base inválido.")
        meta["metas_anuais"] = {}
        for a in anos:
            v, ok = numero(val.get(f"ano_{a}"))
            if not ok:
                erros.append(f"Linha {n}: '{val.get(f'ano_{a}')}' na Meta {a} não é número.")
            meta["metas_anuais"][str(a)] = v
        if meta["codigo"] in codigos_meta:
            erros.append(f"Linha {n}: código de meta {meta['codigo']} repetido.")
            continue
        codigos_meta.add(meta["codigo"])
        d = diretrizes.setdefault(txt["diretriz_codigo"], {"codigo": txt["diretriz_codigo"],
                                                           "descricao": txt["diretriz"], "objetivos": {}})
        if d["descricao"] != txt["diretriz"]:
            avisos.append(f"Linha {n}: diretriz {d['codigo']} com texto diferente do da primeira linha; mantido o primeiro.")
        o = d["objetivos"].setdefault(txt["objetivo_codigo"], {"codigo": txt["objetivo_codigo"],
                                                               "descricao": txt["objetivo"], "metas": []})
        o["metas"].append(meta)
    estrutura = [{**d, "objetivos": list(d["objetivos"].values())} for d in diretrizes.values()]
    return {"diretrizes": estrutura, "erros": erros, "avisos": avisos,
            "totais": {"diretrizes": len(estrutura),
                       "objetivos": sum(len(d["objetivos"]) for d in estrutura),
                       "metas": len(codigos_meta)}}


# ── Estrutura do PMS ─────────────────────────────────────────────────────────

async def limpar_estrutura(db: AsyncSession, plano: PmsPlano) -> None:
    metas = select(PmsMeta.id).where(PmsMeta.plano_id == plano.id)
    await db.execute(delete(PmsMetaAnual).where(PmsMetaAnual.meta_id.in_(metas)))
    await db.execute(delete(PmsMeta).where(PmsMeta.plano_id == plano.id))
    await db.execute(delete(PmsObjetivo).where(PmsObjetivo.plano_id == plano.id))
    await db.execute(delete(PmsDiretriz).where(PmsDiretriz.plano_id == plano.id))


def anos_do(plano: PmsPlano) -> list[int]:
    return list(range(plano.vigencia_inicio, plano.vigencia_fim + 1))


async def gravar_estrutura(db: AsyncSession, plano: PmsPlano, diretrizes: list[dict]) -> dict:
    anos = anos_do(plano)
    tot = {"diretrizes": 0, "objetivos": 0, "metas": 0}
    for i, d in enumerate(diretrizes):
        dr = PmsDiretriz(plano_id=plano.id, municipio_id=plano.municipio_id, codigo=d["codigo"],
                         descricao=d["descricao"], ordem=i)
        db.add(dr)
        await db.flush()
        tot["diretrizes"] += 1
        for j, o in enumerate(d.get("objetivos", [])):
            ob = PmsObjetivo(diretriz_id=dr.id, plano_id=plano.id, municipio_id=plano.municipio_id,
                             codigo=o["codigo"], descricao=o["descricao"], ordem=j)
            db.add(ob)
            await db.flush()
            tot["objetivos"] += 1
            for k, m in enumerate(o.get("metas", [])):
                me = PmsMeta(objetivo_id=ob.id, plano_id=plano.id, municipio_id=plano.municipio_id,
                             uid=m.get("uid") or str(uuid.uuid4()), codigo=m["codigo"], descricao=m["descricao"],
                             ordem=k, **{c: m.get(c) for c in (
                                 "indicador", "tipo_indicador", "unidade", "polaridade", "formula", "linha_base",
                                 "ano_linha_base", "meta_quadrienal", "fonte", "setor_responsavel", "fonte_automatica")})
                db.add(me)
                await db.flush()
                tot["metas"] += 1
                for a in anos:
                    db.add(PmsMetaAnual(meta_id=me.id, municipio_id=plano.municipio_id, ano=a,
                                        meta_prevista=(m.get("metas_anuais") or {}).get(str(a))))
    await db.flush()
    return tot


async def arvore(db: AsyncSession, plano: PmsPlano) -> list[dict]:
    mid = plano.municipio_id
    drs = (await db.execute(select(PmsDiretriz).where(PmsDiretriz.plano_id == plano.id, PmsDiretriz.municipio_id == mid)
                            .order_by(PmsDiretriz.ordem, PmsDiretriz.id))).scalars().all()
    obs = (await db.execute(select(PmsObjetivo).where(PmsObjetivo.plano_id == plano.id, PmsObjetivo.municipio_id == mid)
                            .order_by(PmsObjetivo.ordem, PmsObjetivo.id))).scalars().all()
    mts = (await db.execute(select(PmsMeta).where(PmsMeta.plano_id == plano.id, PmsMeta.municipio_id == mid)
                            .order_by(PmsMeta.ordem, PmsMeta.id))).scalars().all()
    anuais = (await db.execute(select(PmsMetaAnual).where(
        PmsMetaAnual.meta_id.in_([m.id for m in mts] or [0])).order_by(PmsMetaAnual.ano))).scalars().all()
    por_meta: dict[int, list] = {}
    for a in anuais:
        por_meta.setdefault(a.meta_id, []).append({"ano": a.ano, "meta_prevista": a.meta_prevista,
                                                   "resultado": a.resultado})
    n_acoes = {}
    for meta_id, n in (await db.execute(
            select(PmsAcaoMeta.meta_id, PmsAcaoMeta.acao_id).join(PmsAcao, PmsAcao.id == PmsAcaoMeta.acao_id)
            .where(PmsAcaoMeta.municipio_id == mid, PmsAcao.excluido_em.is_(None)))).all():
        n_acoes[meta_id] = n_acoes.get(meta_id, 0) + 1
    return [{"id": d.id, "codigo": d.codigo, "descricao": d.descricao, "objetivos": [
        {"id": o.id, "codigo": o.codigo, "descricao": o.descricao, "metas": [
            {**dict_meta(m), "metas_anuais": por_meta.get(m.id, []), "acoes": n_acoes.get(m.id, 0)}
            for m in mts if m.objetivo_id == o.id]}
        for o in obs if o.diretriz_id == d.id]} for d in drs]


def dict_meta(m: PmsMeta) -> dict:
    return {k: getattr(m, k) for k in (
        "id", "uid", "codigo", "descricao", "indicador", "tipo_indicador", "unidade", "polaridade", "formula",
        "linha_base", "ano_linha_base", "meta_quadrienal", "fonte", "setor_responsavel", "objetivo_id",
        "fonte_automatica")}


def dict_plano(p: PmsPlano) -> dict:
    return {k: (getattr(p, k).isoformat() if isinstance(getattr(p, k), (date, datetime)) else getattr(p, k)) for k in (
        "id", "titulo", "vigencia_inicio", "vigencia_fim", "versao", "situacao", "plano_origem_id", "motivo_versao",
        "resolucao_cms", "aprovado_por", "aprovado_em", "observacoes", "criado_por", "criado_em")}


async def nova_versao(db: AsyncSession, origem: PmsPlano, usuario: str, motivo: str) -> PmsPlano:
    """Proposta de alteração: copia a estrutura (com os mesmos `uid` das metas) para um rascunho."""
    ultima = max((await db.execute(select(PmsPlano.versao).where(
        PmsPlano.municipio_id == origem.municipio_id, PmsPlano.vigencia_inicio == origem.vigencia_inicio))).scalars())
    novo = PmsPlano(municipio_id=origem.municipio_id, titulo=origem.titulo, vigencia_inicio=origem.vigencia_inicio,
                    vigencia_fim=origem.vigencia_fim, versao=ultima + 1, situacao="rascunho",
                    plano_origem_id=origem.id, motivo_versao=motivo, observacoes=origem.observacoes, criado_por=usuario)
    db.add(novo)
    await db.flush()
    estrutura = await arvore(db, origem)
    for d in estrutura:
        for o in d["objetivos"]:
            for m in o["metas"]:
                m["metas_anuais"] = {str(a["ano"]): a["meta_prevista"] for a in m["metas_anuais"]}
    await gravar_estrutura(db, novo, estrutura)
    return novo


# ── Ações e execução ─────────────────────────────────────────────────────────

def acao_atrasada(a: PmsAcao, hoje: date | None = None) -> bool:
    hoje = hoje or date.today()
    return bool(a.prazo and a.prazo < hoje and a.situacao not in SITUACOES_ENCERRADAS)


def prazo_proximo(a: PmsAcao, hoje: date | None = None) -> bool:
    hoje = hoje or date.today()
    return bool(a.prazo and hoje <= a.prazo and (a.prazo - hoje).days <= PRAZO_PROXIMO_DIAS
                and a.situacao not in SITUACOES_ENCERRADAS)


def dict_acao(a: PmsAcao) -> dict:
    d = {k: getattr(a, k) for k in (
        "id", "pas_id", "codigo", "descricao", "produto", "quantidade_prevista", "setor_responsavel",
        "responsavel_usuario_id", "responsavel_nome", "recursos", "fonte_financiamento", "indicador_acompanhamento",
        "evidencia_esperada", "situacao", "percentual_execucao", "quantidade_realizada", "resultado_alcancado",
        "justificativa", "situacao_validada_por")}
    d["prazo"] = a.prazo.isoformat() if a.prazo else None
    d["situacao_validada_em"] = a.situacao_validada_em.isoformat() if a.situacao_validada_em else None
    d["situacao_rotulo"] = SITUACOES_ACAO.get(a.situacao, a.situacao)
    d["atrasada"] = acao_atrasada(a)
    d["prazo_proximo"] = prazo_proximo(a)
    return d


def dict_atualizacao(u: PmsAtualizacao) -> dict:
    return {"id": u.id, "acao_id": u.acao_id, "situacao": u.situacao,
            "situacao_rotulo": SITUACOES_ACAO.get(u.situacao, u.situacao),
            "percentual_execucao": u.percentual_execucao, "quantidade_realizada": u.quantidade_realizada,
            "resultado_alcancado": u.resultado_alcancado, "justificativa": u.justificativa, "status": u.status,
            "informado_por": u.informado_por, "informado_em": u.informado_em.isoformat(),
            "validado_por": u.validado_por, "validado_em": u.validado_em.isoformat() if u.validado_em else None,
            "parecer": u.parecer}


def validar_atualizacao_informada(situacao: str, justificativa: str | None, resultado: str | None,
                                  percentual: float | None) -> None:
    if situacao not in SITUACOES_ACAO:
        raise ValueError("Situação inválida.")
    if situacao in EXIGE_JUSTIFICATIVA and not (justificativa or "").strip():
        raise ValueError(f"Informe a justificativa para \"{SITUACOES_ACAO[situacao]}\".")
    if situacao == "concluida" and not (resultado or "").strip():
        raise ValueError("Informe o resultado alcançado para registrar a conclusão.")
    if percentual is not None and not 0 <= percentual <= 100:
        raise ValueError("Percentual de execução deve estar entre 0 e 100.")


async def tem_evidencia(db: AsyncSession, acao: PmsAcao) -> bool:
    return (await db.execute(select(PmsAnexo.id).where(
        PmsAnexo.municipio_id == acao.municipio_id, PmsAnexo.entidade == "acao", PmsAnexo.entidade_id == acao.id,
        PmsAnexo.excluido_em.is_(None)).limit(1))).first() is not None


async def metas_da_acao(db: AsyncSession, acao: PmsAcao) -> list[dict]:
    rows = (await db.execute(select(PmsMeta).join(PmsAcaoMeta, PmsAcaoMeta.meta_id == PmsMeta.id).where(
        PmsAcaoMeta.acao_id == acao.id, PmsAcaoMeta.municipio_id == acao.municipio_id))).scalars().all()
    return [{"id": m.id, "uid": m.uid, "codigo": m.codigo, "descricao": m.descricao} for m in rows]


async def definir_metas(db: AsyncSession, acao: PmsAcao, plano_id: int, meta_ids: list[int]) -> None:
    validas = set((await db.execute(select(PmsMeta.id).where(
        PmsMeta.id.in_(meta_ids or [0]), PmsMeta.plano_id == plano_id,
        PmsMeta.municipio_id == acao.municipio_id))).scalars())
    if set(meta_ids) - validas:
        raise ValueError("Meta inexistente ou de outro PMS.")
    await db.execute(delete(PmsAcaoMeta).where(PmsAcaoMeta.acao_id == acao.id))
    for mid in sorted(validas):
        db.add(PmsAcaoMeta(acao_id=acao.id, meta_id=mid, municipio_id=acao.municipio_id))
