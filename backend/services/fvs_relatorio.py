"""
Relatórios epidemiológicos automatizados da Vigilância em Saúde (FVS-RCP/AM) — PDF, Excel e Word.

Cada relatório é montado como um documento neutro (seções com parágrafos e tabelas) e depois
desenhado no formato pedido. Fatos vêm das fontes oficiais integradas; a seção "Análise da IA",
quando pedida, é separada e rotulada como interpretação.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date, datetime

from sqlalchemy.ext.asyncio import AsyncSession

from services import fvs_analise as an
from services import fvs_catalogo as cat
from services.fvs_coleta import IBGE_AM

TITULOS = {
    "municipal": "Relatório Epidemiológico Municipal",
    "consolidado": "Relatório Consolidado de Vigilância em Saúde",
    "agravo": "Relatório por Doença ou Agravo",
    "comparativo": "Relatório Comparativo entre Municípios",
    "pqa_vs": "Relatório de Indicadores do PQA-VS",
    "evolucao": "Relatório de Evolução Epidemiológica",
    "alertas": "Relatório de Alertas e Tendências",
    "estadual": "Relatório Estadual Consolidado do Amazonas",
}


class SemDados(RuntimeError):
    pass


@dataclass
class Tabela:
    titulo: str
    colunas: list[str]
    linhas: list[list]
    nota: str | None = None


@dataclass
class Secao:
    titulo: str
    paragrafos: list[str] = field(default_factory=list)
    tabelas: list[Tabela] = field(default_factory=list)
    grafico: dict | None = None       # {"titulo", "rotulos": [...], "valores": [...]}
    rotulo_ia: bool = False


@dataclass
class Documento:
    titulo: str
    municipio: str
    periodo: str
    gerado_por: str
    secoes: list[Secao] = field(default_factory=list)
    fontes: list[str] = field(default_factory=list)
    limitacoes: list[str] = field(default_factory=list)


def _n(v, casas=0):
    if v is None:
        return "—"
    if isinstance(v, (int, float)):
        s = f"{v:,.{casas}f}" if casas else f"{v:,.0f}"
        return s.replace(",", "X").replace(".", ",").replace("X", ".")
    return str(v)


def _pct(v):
    return "—" if v is None else (("+" if v > 0 else "") + _n(v, 1) + "%")


# ── Montagem ──────────────────────────────────────────────────────────────────

async def _fontes(db: AsyncSession, quais=None) -> list[str]:
    out = []
    for f in quais or cat.COLETORES:
        e = await an._estado(db, f)
        quando = f"{e['ultima_sincronizacao']:%d/%m/%Y %H:%M} UTC" if e.get("ultima_sincronizacao") else "nunca sincronizada"
        out.append(f"{e['rotulo']} — {e.get('data_referencia') or 'sem dado'}; publicado pela FVS em "
                   f"{e['publicado_na_fonte']:%d/%m/%Y}" if e.get("publicado_na_fonte") else f"{e['rotulo']} — sem publicação lida")
        out[-1] += f"; sincronizado em {quando}. Painel: {cat.FONTE_PAGINA}"
    return out


def _sec_agravos(p: dict, ano: int) -> Secao:
    linhas = []
    for a in p["agravos"]:
        c = a["comparacao"] or {}
        linhas.append([a["agravo"], _n(a["casos_ano"]), _n(a["casos_ano_anterior"]),
                       f"{_n(c.get('acumulado'))} × {_n(c.get('mesmo_periodo_ano_anterior'))} (SE 1–{c.get('ate_se')})" if c else "—",
                       _pct(c.get("variacao_pct")), _n(a["incidencia_100mil"], 1)])
    pop = p.get("populacao_referencia")
    return Secao("Doenças e agravos de notificação (SINAN — base FVS)",
                 [f"Contagem de notificações por doença/agravo em {ano} (ano parcial) e {ano - 1}. "
                  "A coluna 'mesmo período' compara as mesmas semanas epidemiológicas dos dois anos."],
                 [Tabela(f"Notificações por agravo — {p['municipio']['nome']}",
                         ["Doença/agravo", f"{ano}", f"{ano - 1}", "Mesmo período", "Variação", "Incid./100 mil"], linhas,
                         nota=(f"Incidência calculada pelo ERSUS360: casos ÷ população {pop['ano']} do PQA-VS (FVS) × 100.000."
                               if pop else "Sem população de referência para calcular incidência."))])


def _sec_malaria(p: dict, ano: int) -> Secao | None:
    m = p.get("malaria")
    if not m:
        return None
    ipa = m.get("ipa_ano_anterior") or {}
    c = m.get("comparacao") or {}
    return Secao("Malária (SIVEP-Malária — painel FVS)", [
        f"Exames positivos em {ano}: {_n(m['positivos_ano'])}; em {ano - 1}: {_n(m['positivos_ano_anterior'])}.",
        f"IPA {ano - 1}: {_n(ipa.get('valor'), 2)} por 1.000 habitantes — {ipa.get('texto') or 'sem classificação'} "
        "(positivos ÷ população da planilha da FVS × 1.000; calculado pelo ERSUS360 com a regra do painel).",
        (f"Acumulado até a SE {c['ate_se']}: {_n(c['acumulado'])} contra {_n(c['mesmo_periodo_ano_anterior'])} "
         f"no mesmo período de {ano - 1} ({_pct(c['variacao_pct'])}).") if c else "Sem série semanal.",
    ])


def _sec_pqa(pqa: dict | None) -> Secao:
    if not pqa:
        return Secao("PQA-VS", ["A FVS ainda não publicou resultado do PQA-VS para este município no ERSUS360."])
    linhas = [[i["indicador"].replace("IND_", ""), i["nome"] or "—", _n(i["numerador"], 2 if i["numerador"] and i["numerador"] % 1 else 0),
               _n(i["denominador"], 2 if i["denominador"] and i["denominador"] % 1 else 0), _n(i["resultado"], 1), i["meta"] or "—",
               i["alcance"] or "—", i["observacao"] or ""] for i in pqa["itens"]]
    return Secao(f"PQA-VS — {pqa['periodo_rotulo']}",
                 [f"Metas atingidas: {pqa['atingidos']} de {pqa['avaliados']} indicadores avaliados (alcance publicado pela FVS; "
                  "indicadores com 'ausência de casos' no período não entram na conta)."],
                 [Tabela("Indicadores do PQA-VS", ["Nº", "Indicador", "Num.", "Den.", "Resultado", "Meta", "Atingiu", "Obs."],
                         linhas, nota="De 2024 em diante os indicadores seguem a Portaria GM/MS nº 6.878/2025; a descrição "
                                      "não é exibida porque a pasta pública não a liga ao número do indicador.")])


def _sec_alertas(al: dict) -> Secao:
    linhas = [[{"oficial": "Oficial", "calculado": "Calculado", "integracao": "Integração"}[a["tipo"]], a["nivel"].upper(),
               a["titulo"], a["detalhe"]] for a in al["alertas"]]
    return Secao("Situações que exigem atenção",
                 [al["legenda"]["oficial"], al["legenda"]["calculado"]] if linhas else ["Nenhum alerta no período."],
                 [Tabela("Alertas", ["Tipo", "Nível", "Alerta", "Detalhe"], linhas)] if linhas else [])


async def _evolucao(db: AsyncSession, m: str, ano: int, nomes: dict) -> Secao:
    rows = await an.serie(db, "sinan", None, [m], ano - 9, ano, "ano")
    anos = list(range(ano - 9, ano + 1))
    por: dict[str, dict[int, float]] = {}
    for r in rows:
        por.setdefault(r["indicador"][6:], {})[r["ano"]] = r["valor"]
    top = sorted(por, key=lambda k: -sum(por[k].values()))[:12]
    linhas = [[ag] + [_n(por[ag].get(a)) for a in anos] for ag in top]
    total = [sum(v.get(a, 0) for v in por.values()) for a in anos]
    return Secao("Evolução temporal — notificações por ano",
                 [f"Série {anos[0]}–{ano} ({ano} parcial). Principais agravos por volume em {nomes.get(m, m)}."],
                 [Tabela("Notificações por ano", ["Agravo"] + [str(a) for a in anos], linhas)],
                 grafico={"titulo": "Total de notificações (todas as doenças/agravos)", "rotulos": [str(a) for a in anos],
                          "valores": total})


async def montar(db: AsyncSession, tipo: str, ibges: list[str], ano: int, agravo: str | None,
                 gerado_por: str) -> Documento:
    nomes = await an.nomes_municipios(db)
    if len(nomes) <= 1:
        raise SemDados("Os dados da FVS ainda não foram sincronizados no ERSUS360.")
    m = ibges[0] if tipo != "estadual" else IBGE_AM
    doc = Documento(TITULOS[tipo], nomes.get(m, m) if tipo != "comparativo" else ", ".join(nomes.get(x, x) for x in ibges),
                    f"{ano} (até a data de referência de cada fonte) e série histórica", gerado_por)
    doc.limitacoes = [
        "A base pública de monitoramento SINAN da FVS traz contagens de notificação, sem classificação final "
        "(confirmado/descartado/em investigação) nem óbitos.",
        "Painéis Power BI da FVS (mortalidade materna e infantil, nascimentos, SRAG, Mpox, esporotricose, PES, "
        "tratamento preventivo da TB, população negra) não oferecem exportação de dados: consulte-os nas subabas.",
        "Painéis de IST, sífilis, HIV/aids, hepatites, tuberculose e óbitos de mulheres não são extraídos: as pastas "
        "públicas contêm dados pessoais identificados (proteção de dados — LGPD).",
        "Dados do ano corrente são parciais e podem ser revisados pela FVS a cada publicação.",
    ]
    if tipo in {"municipal", "consolidado", "alertas"}:
        p = await an.painel_executivo(db, m, ano)
        c = p["cartoes"]
        doc.secoes.append(Secao("Resumo executivo", [
            f"Notificações SINAN em {ano}: {_n(c['notificacoes_sinan']['valor'])} (todas as doenças/agravos da base FVS).",
            f"Malária — exames positivos em {ano}: {_n(c['malaria_positivos']['valor'])}.",
            f"Casos de SRAG (SIVEP-Gripe, painel AESOP) em {ano}: {_n(c['srag_casos']['valor'])}.",
            (f"PQA-VS ({c['pqa_vs']['periodo']}): {c['pqa_vs']['valor']} de {c['pqa_vs']['total']} metas atingidas."
             if c["pqa_vs"]["valor"] is not None else "PQA-VS: sem resultado publicado."),
        ]))
        if tipo != "alertas":
            doc.secoes.append(_sec_agravos(p, ano))
            if (s := _sec_malaria(p, ano)):
                doc.secoes.append(s)
            if p.get("aesop"):
                a = p["aesop"]
                doc.secoes.append(Secao(f"Síndromes respiratórias — AESOP (SE {a['se']}/{a['ano']})", [
                    f"Aviso AESOP: {a.get('aviso_aesop', '—')}; Aviso APS: {a.get('aviso_aps', '—')}; "
                    f"Aviso PIP: {a.get('aviso_pip', '—')}; probabilidade de crescimento: {_n(a.get('prob_crescimento_pct'), 1)}%; "
                    f"atendimentos de síndrome gripal na APS: {_n(a.get('sindrome_gripal_aps'))} de "
                    f"{_n(a.get('atendimentos_aps'))} atendimentos (publicado pelo CIEVS/FVS)."]))
            doc.secoes.append(_sec_pqa(p["pqa_vs"]))
        doc.secoes.append(_sec_alertas(await an.alertas(db, m, ano)))
        if tipo == "consolidado":
            doc.secoes.append(await _evolucao(db, m, ano, nomes))
    elif tipo == "evolucao":
        doc.secoes.append(await _evolucao(db, m, ano, nomes))
        mal = await an.serie(db, "malaria", None, [m], ano - 9, ano, "ano")
        if mal:
            anos = sorted({r["ano"] for r in mal})
            d = {(r["indicador"], r["ano"]): r for r in mal}
            doc.secoes.append(Secao("Malária — positivos e IPA por ano", [], [Tabela(
                "Malária", ["Ano", "Positivos", "IPA", "Classe"],
                [[a, _n((d.get(("positivos", a)) or {}).get("valor")), _n((d.get(("ipa", a)) or {}).get("valor"), 2),
                  (d.get(("ipa", a)) or {}).get("texto") or "—"] for a in anos])]))
    elif tipo == "agravo":
        if not agravo:
            raise SemDados("Escolha a doença/agravo do relatório.")
        ind = f"casos:{agravo}"
        anual = await an.serie(db, "sinan", ind, [m, IBGE_AM], ano - 9, ano, "ano")
        if not anual:
            raise SemDados(f"Sem notificações de {agravo} na base da FVS para o período.")
        anos = list(range(ano - 9, ano + 1))
        d = {(r["ibge6"], r["ano"]): r["valor"] for r in anual}
        comp = await an.comparar_periodo(db, "sinan", ind, m, ano)
        doc.titulo = f"{TITULOS['agravo']} — {agravo}"
        doc.secoes.append(Secao("Resumo", [
            f"{agravo} em {nomes.get(m)}: {_n(d.get((m, ano)))} notificações em {ano} e {_n(d.get((m, ano - 1)))} em {ano - 1}.",
            (f"Até a SE {comp['ate_se']}: {_n(comp['acumulado'])} × {_n(comp['mesmo_periodo_ano_anterior'])} no mesmo período "
             f"do ano anterior ({_pct(comp['variacao_pct'])}); últimas {comp['janela_semanas']} semanas: "
             f"{_n(comp['ultimas_semanas'])} × {_n(comp['semanas_anteriores'])} ({_pct(comp['variacao_recente_pct'])}).")
            if comp else "Sem série semanal no ano."],
            [Tabela("Notificações por ano", ["Ano", nomes.get(m, m), "Amazonas"],
                    [[a, _n(d.get((m, a))), _n(d.get((IBGE_AM, a)))] for a in anos])],
            grafico={"titulo": f"{agravo} — {nomes.get(m)}", "rotulos": [str(a) for a in anos],
                     "valores": [d.get((m, a)) or 0 for a in anos]}))
        sem = await an.serie(db, "sinan", ind, [m], ano, ano, "se")
        if sem:
            doc.secoes.append(Secao(f"Semanas epidemiológicas de {ano}", [], [Tabela(
                "Notificações por SE", ["SE", "Casos"], [[r["periodo"], _n(r["valor"])] for r in sem])]))
        for rec, rot in (("sexo", "Sexo"), ("faixa_etaria", "Faixa etária")):
            rr = await an.serie(db, "sinan", ind, [m], ano - 1, ano, "ano", rec)
            if rr:
                doc.secoes.append(Secao(f"Perfil por {rot.lower()}", [], [Tabela(
                    rot, ["Ano", rot, "Casos"], [[r["ano"], r["recorte_valor"], _n(r["valor"])] for r in rr])]))
        doc.secoes.append(Secao("Municípios do AM com maior aumento", [], [Tabela(
            "Tendência", ["Município", "Acumulado", "Mesmo período ano ant.", "Variação"],
            [[t["nome"], _n(t["acumulado"]), _n(t["mesmo_periodo_ano_anterior"]), _pct(t["variacao_pct"])]
             for t in await an.tendencias(db, "sinan", ind, ano, 10)])]))
    elif tipo == "comparativo":
        comp = await an.comparativo(db, ibges, ano)
        cols = ["Indicador"] + [nomes.get(x, x) for x in ibges]
        linhas = []
        for ln in comp["linhas"]:
            linhas.append([("Malária: " if ln["fonte"] == "malaria" else "") + ln["indicador"].replace("casos:", "")] + [
                (_n(ln["valores"][x]["valor"], 2 if ln["indicador"] == "ipa" else 0)
                 + (f" ({_n(ln['valores'][x]['taxa_100mil'], 1)}/100 mil)" if ln["valores"].get(x, {}).get("taxa_100mil") else ""))
                if x in ln["valores"] else "—" for x in ibges])
        doc.secoes.append(Secao(f"Comparação em {ano}", ["Taxas por 100 mil calculadas pelo ERSUS360 com a população do PQA-VS (FVS)."],
                                [Tabela("Indicadores por município", cols, linhas)]))
        doc.secoes.append(Secao("PQA-VS", [f"{nomes.get(x, x)}: " + (f"{v['atingidos']} de {v['avaliados']} metas ({v['periodo']})" if v else "sem resultado")
                                           for x, v in comp["pqa_vs"].items()]))
    elif tipo == "pqa_vs":
        doc.secoes.append(_sec_pqa(await an.pqa_resumo(db, m)))
        hist = await an.pqa_historico(db, m)
        doc.secoes.append(Secao("Histórico", ["Resultados publicados pela FVS por ano (até 2023) e por quadrimestre (2025 em diante)."],
                                [Tabela("Histórico PQA-VS", ["Período", "Nº", "Resultado", "Meta", "Atingiu", "Obs."],
                                        [[h["periodo"], h["indicador"][4:], _n(h["resultado"], 1), h["meta"] or "—",
                                          h["alcance"] or "—", h["observacao"] or ""] for h in hist])]))
    elif tipo == "estadual":
        p = await an.painel_executivo(db, IBGE_AM, ano)
        doc.secoes.append(_sec_agravos(p, ano))
        if (s := _sec_malaria(p, ano)):
            doc.secoes.append(s)
        for ind in [a["indicador"] for a in p["agravos"][:5]]:
            mp = await an.mapa(db, "sinan", ind, ano)
            top = sorted([x for x in mp["municipios"] if x["valor"]], key=lambda x: -x["valor"])[:10]
            doc.secoes.append(Secao(f"{ind[6:]} — municípios com mais notificações em {ano}", [], [Tabela(
                ind[6:], ["Município", "Casos", "Por 100 mil", f"{ano - 1}"],
                [[x["nome"], _n(x["valor"]), _n(x["taxa_100mil"], 1), _n(x["ano_anterior"])] for x in top])]))
        ipa = await an.mapa(db, "malaria", "ipa", ano - 1)
        top = sorted([x for x in ipa["municipios"] if x["valor"]], key=lambda x: -x["valor"])[:15]
        doc.secoes.append(Secao(f"Malária — IPA {ano - 1} por município (maiores)", [], [Tabela(
            "IPA", ["Município", "IPA", "Classe"], [[x["nome"], _n(x["valor"], 2), x["texto"]] for x in top])]))
    doc.fontes = await _fontes(db)
    if not any(s.tabelas or s.paragrafos for s in doc.secoes):
        raise SemDados("Sem dados para o relatório pedido.")
    return doc


# ── Desenho ───────────────────────────────────────────────────────────────────

def _pdf(doc: Documento) -> bytes:
    from reportlab.graphics.charts.barcharts import VerticalBarChart
    from reportlab.graphics.shapes import Drawing
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape

    buf = io.BytesIO()
    larga = any(len(t.colunas) > 6 for s in doc.secoes for t in s.tabelas)
    tam = landscape(A4) if larga else A4
    pdf = SimpleDocTemplate(buf, pagesize=tam, leftMargin=1.5 * cm, rightMargin=1.5 * cm, topMargin=1.5 * cm,
                            bottomMargin=1.5 * cm, title=doc.titulo, author="ERSUS360")
    ss = getSampleStyleSheet()
    cel = ParagraphStyle("cel", parent=ss["BodyText"], fontSize=7.5, leading=9)
    ia = ParagraphStyle("ia", parent=ss["BodyText"], backColor=colors.HexColor("#f1f5ff"), borderPadding=4)
    el = [Paragraph(escape(doc.titulo), ss["Title"]),
          Paragraph(escape(f"{doc.municipio} — Amazonas · Período: {doc.periodo}"), ss["Normal"]),
          Paragraph(escape(f"Gerado em {datetime.now():%d/%m/%Y %H:%M} por {doc.gerado_por} · ERSUS360"), ss["Normal"]),
          Spacer(1, 10)]
    largura = tam[0] - 3 * cm
    for s in doc.secoes:
        el.append(Paragraph(escape(s.titulo) + (" <font color='#1d4ed8'>(interpretação gerada por IA)</font>" if s.rotulo_ia else ""), ss["Heading2"]))
        for p in s.paragrafos:
            el.append(Paragraph(escape(p).replace("\n", "<br/>"), ia if s.rotulo_ia else ss["BodyText"]))
        if s.grafico and any(s.grafico["valores"]):
            d = Drawing(largura, 160)
            ch = VerticalBarChart()
            ch.x, ch.y, ch.width, ch.height = 30, 20, largura - 60, 120
            ch.data = [[float(v or 0) for v in s.grafico["valores"]]]
            ch.categoryAxis.categoryNames = s.grafico["rotulos"]
            ch.bars[0].fillColor = colors.HexColor("#1d4ed8")
            ch.valueAxis.valueMin = 0
            d.add(ch)
            el += [Paragraph(escape(s.grafico["titulo"]), ss["Italic"]), d]
        for t in s.tabelas:
            dados = [[Paragraph(f"<b>{escape(str(c))}</b>", cel) for c in t.colunas]] + \
                    [[Paragraph(escape(str(v)), cel) for v in ln] for ln in t.linhas]
            tb = Table(dados, repeatRows=1, colWidths=[largura / len(t.colunas)] * len(t.colunas))
            tb.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e2e8f0")),
                                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#94a3b8")),
                                    ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            el += [Spacer(1, 4), tb]
            if t.nota:
                el.append(Paragraph(escape(t.nota), ParagraphStyle("n", parent=cel, textColor=colors.HexColor("#475569"))))
        el.append(Spacer(1, 8))
    el.append(Paragraph("Fontes oficiais", ss["Heading2"]))
    el += [Paragraph(escape("• " + f), cel) for f in doc.fontes]
    el.append(Paragraph("Limitações e informações indisponíveis", ss["Heading2"]))
    el += [Paragraph(escape("• " + f), cel) for f in doc.limitacoes]
    pdf.build(el)
    return buf.getvalue()


def _xlsx(doc: Documento) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    ws.append([doc.titulo])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"{doc.municipio} — Amazonas"])
    ws.append([f"Período: {doc.periodo}"])
    ws.append([f"Gerado em {datetime.now():%d/%m/%Y %H:%M} por {doc.gerado_por} · ERSUS360"])
    for s in doc.secoes:
        ws.append([])
        ws.append([s.titulo + (" (interpretação gerada por IA)" if s.rotulo_ia else "")])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        for p in s.paragrafos:
            ws.append([p])
    ws.append([])
    ws.append(["Fontes oficiais"])
    for f in doc.fontes:
        ws.append([f])
    ws.append(["Limitações"])
    for f in doc.limitacoes:
        ws.append([f])
    ws.column_dimensions["A"].width = 120
    usados = set()
    for s in doc.secoes:
        for t in s.tabelas:
            nome = "".join(ch for ch in t.titulo if ch not in "[]:*?/\\")[:28] or "Tabela"
            base, i = nome, 2
            while nome in usados:
                nome, i = f"{base[:25]} {i}", i + 1
            usados.add(nome)
            w = wb.create_sheet(nome)
            w.append(t.colunas)
            for c in w[1]:
                c.font = Font(bold=True)
                c.fill = PatternFill("solid", fgColor="E2E8F0")
            for ln in t.linhas:
                w.append([v for v in ln])
            if t.nota:
                w.append([])
                w.append([t.nota])
            for col in w.columns:
                w.column_dimensions[col[0].column_letter].width = 18
            w.column_dimensions["A"].width = 40
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _docx(doc: Documento) -> bytes:
    from docx import Document as Docx
    from docx.shared import Pt

    d = Docx()
    d.add_heading(doc.titulo, 0)
    d.add_paragraph(f"{doc.municipio} — Amazonas · Período: {doc.periodo}")
    d.add_paragraph(f"Gerado em {datetime.now():%d/%m/%Y %H:%M} por {doc.gerado_por} · ERSUS360")
    for s in doc.secoes:
        d.add_heading(s.titulo + (" (interpretação gerada por IA)" if s.rotulo_ia else ""), 1)
        for p in s.paragrafos:
            d.add_paragraph(p)
        if s.grafico:
            d.add_paragraph(s.grafico["titulo"] + ": " + "; ".join(
                f"{r}: {_n(v)}" for r, v in zip(s.grafico["rotulos"], s.grafico["valores"])))
        for t in s.tabelas:
            d.add_paragraph(t.titulo).runs[0].bold = True
            tb = d.add_table(rows=1, cols=len(t.colunas))
            tb.style = "Light Grid Accent 1"
            for i, c in enumerate(t.colunas):
                tb.rows[0].cells[i].text = str(c)
            for ln in t.linhas:
                cells = tb.add_row().cells
                for i, v in enumerate(ln):
                    cells[i].text = str(v)
            for row in tb.rows:
                for c in row.cells:
                    for p in c.paragraphs:
                        for r in p.runs:
                            r.font.size = Pt(8)
            if t.nota:
                d.add_paragraph(t.nota).runs[0].italic = True
    d.add_heading("Fontes oficiais", 1)
    for f in doc.fontes:
        d.add_paragraph(f, style="List Bullet")
    d.add_heading("Limitações e informações indisponíveis", 1)
    for f in doc.limitacoes:
        d.add_paragraph(f, style="List Bullet")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


MIME = {"pdf": "application/pdf",
        "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}


async def gerar(db: AsyncSession, tipo: str, formato: str, ibges: list[str], ano: int, agravo: str | None,
                incluir_ia: bool, gerado_por: str) -> tuple[bytes, str, str]:
    doc = await montar(db, tipo, ibges, ano, agravo, gerado_por)
    if incluir_ia:
        from services import fvs_ia
        try:
            r = await fvs_ia.responder(db, f"Redija as considerações técnicas para a gestão do {doc.titulo} de "
                                           f"{doc.municipio}, ano {ano}: tendências, situações que exigem atenção e "
                                           "recomendações. Use as ferramentas para conferir os números.", ibges[0])
            doc.secoes.append(Secao("Considerações técnicas para a gestão", [r["resposta"]], rotulo_ia=True))
        except fvs_ia.IAIndisponivel as exc:
            doc.secoes.append(Secao("Considerações técnicas para a gestão", [f"Análise por IA indisponível: {exc}"], rotulo_ia=True))
    render = {"pdf": _pdf, "xlsx": _xlsx, "docx": _docx}[formato]
    slug = "".join(ch if ch.isalnum() else "-" for ch in f"{tipo}-{doc.municipio}-{ano}".lower())[:80]
    return render(doc), MIME[formato], f"vigilancia-{slug}.{formato}"
