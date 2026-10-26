"""
PMS Vivo — relatórios e exportações (PDF e Excel). Cada relatório é montado como documento neutro
(seções com parágrafos e tabelas) e depois desenhado no formato pedido. Só usa o que está registrado:
campo vazio aparece como "—", nunca estimado.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import datetime

from services import pms_ciclo as ciclo
from services import pms_monitoramento as mon
from services import pms_vivo as base

TIPOS = {
    "execucao_pas": "Relatório de Execução da PAS",
    "metas_pms": "Relatório de Metas do PMS",
    "pendencias": "Relatório de Pendências",
    "por_responsavel": "Relatório por Responsável",
    "quadrimestral": "Relatório Quadrimestral (apoio ao RDQA)",
    "anual": "Avaliação Anual (apoio ao RAG)",
    "reprogramacoes": "Relatório de Reprogramações",
    "plano_acao": "Plano de Ação 5W2H (ações corretivas)",
}


@dataclass
class Secao:
    titulo: str
    paragrafos: list[str] = field(default_factory=list)
    tabelas: list[tuple[list[str], list[list]]] = field(default_factory=list)


@dataclass
class Documento:
    titulo: str
    subtitulo: str
    secoes: list[Secao] = field(default_factory=list)
    aviso: str | None = None


def _v(x) -> str:
    if x is None or x == "":
        return "—"
    if isinstance(x, float):
        return f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".").rstrip("0").rstrip(",")
    return str(x)


def _br(d: str | None) -> str:
    return "/".join(reversed(d[:10].split("-"))) if d else "—"


def _sub(municipio: str, plano: dict, extra: str) -> str:
    return f"{municipio} · {plano['titulo']} (v{plano['versao']}, {plano['situacao']}) · {extra}"


async def montar(db, tipo: str, plano, municipio: str, ano: int, quadrimestre: int | None = None,
                 responsavel: str | None = None) -> Documento:
    p = base.dict_plano(plano)
    if tipo == "quadrimestral":
        c = await ciclo.consolidacao_quadrimestral(db, plano, ano, quadrimestre or 1)
        doc = Documento(TIPOS[tipo], _sub(municipio, p, f"{c['quadrimestre']}º quadrimestre de {ano} "
                        f"({_br(c['periodo']['inicio'])} a {_br(c['periodo']['fim'])})"), aviso=c["aviso"])
        doc.secoes += [
            Secao("Ações realizadas no período (validadas)", [] if c["acoes_realizadas"] else ["Nenhuma."],
                  [(["Ação", "Descrição", "Situação", "Resultado"], [[a["codigo"], a["descricao"], a["situacao_rotulo"],
                                                                     _v(a["resultado_alcancado"])] for a in c["acoes_realizadas"]])]
                  if c["acoes_realizadas"] else []),
            Secao("Ações pendentes (prazo até o fim do período)", [] if c["acoes_pendentes"] else ["Nenhuma."],
                  [(["Ação", "Descrição", "Prazo", "Responsável", "Situação"], [[a["codigo"], a["descricao"], _br(a["prazo"]),
                                                                              _v(a["responsavel_nome"]), a["situacao_rotulo"]]
                                                                             for a in c["acoes_pendentes"]])]
                  if c["acoes_pendentes"] else []),
            _secao_metas(c["metas"], ano, "Metas e resultados dos indicadores"),
            Secao("Dificuldades identificadas (FCA)", [] if c["dificuldades"] else ["Nenhuma registrada."],
                  [(["Fato", "Causa", "Causa"], [[d["fato"], d["causa"], "validada" if d["causa_validacao"] == "confirmada"
                                                  else "hipótese"] for d in c["dificuldades"]])] if c["dificuldades"] else []),
            Secao("Providências adotadas", [] if c["providencias"] else ["Nenhuma registrada."],
                  [(["Ação corretiva", "Responsável", "Prazo", "Situação"], [[x["acao_corretiva"], _v(x["responsavel"]),
                                                                           _br(x["prazo"]), x["situacao"]] for x in c["providencias"]])]
                  if c["providencias"] else []),
            Secao("Recomendações", c["recomendacoes"] or ["Nenhuma."]),
        ]
        return doc
    if tipo == "anual":
        a = await ciclo.avaliacao_anual(db, plano, ano)
        doc = Documento(TIPOS[tipo], _sub(municipio, p, f"Exercício {ano}"), aviso=a["aviso"])
        ac, me = a["acoes"], a["metas"]
        doc.secoes += [
            Secao("Síntese", [f"Ações previstas: {ac['previstas']} · executadas: {ac['executadas']} · parcialmente: "
                              f"{ac['parcialmente_executadas']} · não executadas: {ac['nao_executadas']} · canceladas: "
                              f"{ac['canceladas']} · em aberto: {ac['em_aberto']}.",
                              f"Metas: {me['total']} · alcançadas: {me['alcancadas']} · não alcançadas: {me['nao_alcancadas']}"
                              f" · em risco: {me['em_risco']} · sem resultado registrado: {me['sem_resultado']}."]),
            _secao_metas(a["lista_metas"], ano, "Metas do exercício"),
            Secao("Principais resultados", [] if a["principais_resultados"] else ["Nenhum resultado registrado."],
                  [(["Ação", "Descrição", "Resultado"], [[x["acao"], x["descricao"], x["resultado"]]
                                                         for x in a["principais_resultados"]])] if a["principais_resultados"] else []),
            Secao("Justificativas", [] if a["justificativas"] else ["Nenhuma."],
                  [(["Ação", "Situação", "Justificativa"], [[x["acao"], x["situacao"], _v(x["justificativa"])]
                                                            for x in a["justificativas"]])] if a["justificativas"] else []),
            Secao("Dificuldades (FCA)", [] if a["dificuldades"] else ["Nenhuma registrada."],
                  [(["Fato", "Causa", "Causa"], [[d["fato"], d["causa"], "validada" if d["causa_validacao"] == "confirmada"
                                                  else "hipótese"] for d in a["dificuldades"]])] if a["dificuldades"] else []),
            Secao("Ações corretivas", [] if a["acoes_corretivas"] else ["Nenhuma."],
                  [(["Ação corretiva", "Responsável", "Prazo", "Situação"], [[x["acao_corretiva"], _v(x["responsavel"]),
                                                                           _br(x["prazo"]), x["situacao"]]
                                                                          for x in a["acoes_corretivas"]])] if a["acoes_corretivas"] else []),
            _secao_reprog(a["reprogramacoes"], "Propostas de reprogramação"),
            Secao("Recomendações para a PAS seguinte", a["recomendacoes"] or ["Nenhuma."]),
        ]
        return doc
    if tipo == "reprogramacoes":
        doc = Documento(TIPOS[tipo], _sub(municipio, p, f"Ações de {ano}"))
        doc.secoes.append(_secao_reprog(await ciclo.reprogramacoes(db, plano, ano), "Reprogramações"))
        return doc

    metas = await mon.monitoramento(db, plano, ano)
    acoes = await mon.acoes_do_ano(db, plano, ano)
    if tipo == "metas_pms":
        doc = Documento(TIPOS[tipo], _sub(municipio, p, f"situação no ano {ano}"))
        doc.secoes.append(_secao_metas(metas, ano, "Metas: previsto × resultado"))
        return doc
    if tipo == "execucao_pas":
        doc = Documento(TIPOS[tipo], _sub(municipio, p, f"PAS {ano}"))
        doc.secoes.append(Secao("Ações da PAS", [] if acoes else ["Nenhuma ação na PAS do ano."], [(
            ["Ação", "Descrição", "Metas", "Responsável", "Prazo", "Situação oficial", "% exec.", "Resultado"],
            [[a["codigo"], a["descricao"], ", ".join(m["codigo"] for m in a["metas"]), _v(a["responsavel_nome"]),
              _br(a["prazo"]), a["situacao_rotulo"] + (" (atrasada)" if a["atrasada"] else ""),
              _v(a["percentual_execucao"]), _v(a["resultado_alcancado"])] for a in acoes])] if acoes else []))
        return doc
    pend = []
    from sqlalchemy import select
    from models.pms_vivo import PmsAtualizacao
    pend = (await db.execute(select(PmsAtualizacao).where(
        PmsAtualizacao.status == "pendente", PmsAtualizacao.acao_id.in_([a["id"] for a in acoes] or [0])))).scalars().all()
    als = mon.alertas(acoes, metas, pend)
    if tipo == "pendencias":
        doc = Documento(TIPOS[tipo], _sub(municipio, p, f"ano {ano}"))
        doc.secoes.append(Secao("Pendências e alertas", [] if als else ["Nenhuma pendência."],
                                [(["Gravidade", "Pendência", "Responsável"], [[x["severidade"], x["texto"], _v(x["responsavel"])]
                                                                             for x in als])] if als else []))
        return doc
    if tipo == "por_responsavel":
        doc = Documento(TIPOS[tipo], _sub(municipio, p, f"PAS {ano}" + (f" · {responsavel}" if responsavel else "")))
        grupos: dict[str, list] = {}
        for a in acoes:
            if responsavel and a["responsavel_nome"] != responsavel:
                continue
            grupos.setdefault(a["responsavel_nome"] or "Sem responsável", []).append(a)
        for nome, lst in sorted(grupos.items()):
            doc.secoes.append(Secao(nome, [f"{len(lst)} ação(ões) · atrasadas: {sum(a['atrasada'] for a in lst)} · "
                                           f"aguardando validação: {sum(a['aguardando_validacao'] for a in lst)}"],
                                    [(["Ação", "Descrição", "Prazo", "Situação", "Evidência"],
                                      [[a["codigo"], a["descricao"], _br(a["prazo"]), a["situacao_rotulo"]
                                        + (" (atrasada)" if a["atrasada"] else ""),
                                        "anexada" if a["tem_evidencia"] else ("pendente" if a["evidencia_esperada"] else "—")]
                                       for a in lst])]))
        if not grupos:
            doc.secoes.append(Secao("Ações", ["Nenhuma ação."]))
        return doc
    if tipo == "plano_acao":
        fcas = await ciclo._fcas(db, plano, ano)
        doc = Documento(TIPOS[tipo], _sub(municipio, p, f"ano {ano}"),
                        aviso="Ações corretivas registradas nas análises FCA. Onde, como e quanto ficam em branco quando não informados.")
        doc.secoes.append(Secao("5W2H", [] if fcas else ["Nenhuma ação corretiva registrada."], [(
            ["O quê", "Por quê", "Onde", "Quando", "Quem", "Como", "Quanto", "Situação"],
            [[f["acao_corretiva"], f"{f['fato']} — causa ({'validada' if f['causa_validacao'] == 'confirmada' else 'hipótese'}): "
              f"{f['causa']}", "—", _br(f["prazo"]), _v(f["responsavel"]), "—", "—", f["situacao"]] for f in fcas])] if fcas else []))
        return doc
    raise ValueError("Tipo de relatório desconhecido.")


def _secao_metas(metas: list[dict], ano: int, titulo: str) -> Secao:
    if not metas:
        return Secao(titulo, ["Nenhuma meta."])
    anos = [a["ano"] for a in metas[0]["anos"]]
    linhas = []
    for m in metas:
        linhas.append([m["codigo"], m["descricao"]] + [f"{_v(a['meta_prevista'])} → {_v(a['resultado'])}" for a in m["anos"]]
                      + [mon.SITUACOES_META.get(m["situacao"], "—") if m.get("situacao") else "—",
                         f"{_v(m['cumprimento'])}%" if m.get("cumprimento") is not None else "—"])
    return Secao(titulo, [f"Meta prevista → resultado registrado (com fonte). Situação e cumprimento referentes a {ano}."],
                 [(["Meta", "Descrição"] + [str(a) for a in anos] + ["Situação", "Cumprimento"], linhas)])


def _secao_reprog(reprogs: list[dict], titulo: str) -> Secao:
    if not reprogs:
        return Secao(titulo, ["Nenhuma reprogramação."])
    return Secao(titulo, [], [(["Ação", "Situação em origem", "Destino", "Motivo", "Novo prazo", "Nova qtd.", "Decisão"],
                               [[r["acao"]["codigo"], r["situacao_origem"], str(r["ano_destino"]), r["motivo"], _br(r["novo_prazo"]),
                                 _v(r["nova_quantidade"]), r["situacao_rotulo"] + (f" ({r['decidido_por']})" if r["decidido_por"] else "")]
                                for r in reprogs])])


def documento_reuniao(municipio: str, plano: dict, r: dict) -> Documento:
    pa = r["pauta"]
    doc = Documento(f"Reunião de monitoramento do PMS — {r['titulo']}",
                    _sub(municipio, plano, f"{_br(r['data'])} · {r['situacao']}"))
    doc.secoes += [
        Secao("Participantes", [r["participantes"] or "Não registrados."]),
        Secao("Pauta — metas em risco", [] if pa["metas_em_risco"] else ["Nenhuma."],
              [(["Meta", "Descrição", "Situação"], [[m["codigo"], m["descricao"], mon.SITUACOES_META.get(m["situacao"], m["situacao"])]
                                                    for m in pa["metas_em_risco"]])] if pa["metas_em_risco"] else []),
        Secao("Pauta — ações atrasadas", [] if pa["acoes_atrasadas"] else ["Nenhuma."],
              [(["Ação", "Descrição", "Prazo", "Responsável"], [[a["codigo"], a["descricao"], _br(a["prazo"]), _v(a["responsavel"])]
                                                               for a in pa["acoes_atrasadas"]])] if pa["acoes_atrasadas"] else []),
        Secao("Pauta — pendências por responsável", [f"{x['responsavel']}: " + "; ".join(x["itens"])
                                                     for x in pa["pendencias_por_responsavel"]] or ["Nenhuma."]),
        Secao("Pauta — problemas que exigem decisão", pa["problemas_para_decisao"] or ["Nenhum."]),
        Secao("Pauta — reprogramações propostas", [f"{x['acao']} → {x['ano_destino']}: {x['motivo']}"
                                                   for x in pa["reprogramacoes_propostas"]] or ["Nenhuma."]),
        Secao("Pauta — providências da reunião anterior", [f"{x['reuniao']}: {x['decisao']} ({_v(x['responsavel'])}, prazo "
                                                           f"{_br(x['prazo'])})" for x in pa["providencias_anteriores"]] or ["Nenhuma pendente."]),
        Secao("Registro (ata)", [r["ata"] or "Não registrada."]),
        Secao("Decisões", [] if r["decisoes"] else ["Nenhuma."],
              [(["Decisão", "Responsável", "Prazo", "Situação"], [[d["descricao"], _v(d["responsavel"]), _br(d["prazo"]),
                                                                  d["situacao"] + (" (atrasada)" if d["atrasada"] else "")]
                                                                 for d in r["decisoes"]])] if r["decisoes"] else []),
    ]
    return doc


# ── Desenho ───────────────────────────────────────────────────────────────────

def pdf(doc: Documento, gerado_por: str) -> bytes:
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = io.BytesIO()
    tam = landscape(A4)
    pdf_ = SimpleDocTemplate(buf, pagesize=tam, leftMargin=1.5 * cm, rightMargin=1.5 * cm, topMargin=1.5 * cm,
                             bottomMargin=1.5 * cm, title=doc.titulo, author="ERSUS360")
    ss = getSampleStyleSheet()
    cel = ParagraphStyle("cel", parent=ss["BodyText"], fontSize=7.5, leading=9)
    aviso = ParagraphStyle("av", parent=ss["BodyText"], textColor=colors.HexColor("#92400e"))
    largura = tam[0] - 3 * cm
    el = [Paragraph(escape(doc.titulo), ss["Title"]), Paragraph(escape(doc.subtitulo), ss["Normal"]),
          Paragraph(escape(f"Gerado em {datetime.now():%d/%m/%Y %H:%M} por {gerado_por} · ERSUS360 — PMS Vivo"), ss["Normal"])]
    if doc.aviso:
        el.append(Paragraph(escape(doc.aviso), aviso))
    el.append(Spacer(1, 8))
    for s in doc.secoes:
        el.append(Paragraph(escape(s.titulo), ss["Heading2"]))
        for p in s.paragrafos:
            el.append(Paragraph(escape(p), ss["BodyText"]))
        for colunas, linhas in s.tabelas:
            dados = [[Paragraph(f"<b>{escape(str(c))}</b>", cel) for c in colunas]] + \
                    [[Paragraph(escape(_v(v)), cel) for v in ln] for ln in linhas]
            tb = Table(dados, repeatRows=1, colWidths=[largura / len(colunas)] * len(colunas))
            tb.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e2e8f0")),
                                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#94a3b8")),
                                    ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            el += [Spacer(1, 4), tb]
        el.append(Spacer(1, 6))
    pdf_.build(el)
    return buf.getvalue()


def xlsx(doc: Documento, gerado_por: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Relatório"
    ws.append([doc.titulo])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([doc.subtitulo])
    ws.append([f"Gerado em {datetime.now():%d/%m/%Y %H:%M} por {gerado_por} · ERSUS360 — PMS Vivo"])
    if doc.aviso:
        ws.append([doc.aviso])
    for s in doc.secoes:
        ws.append([])
        ws.append([s.titulo])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        for p in s.paragrafos:
            ws.append([p])
        for colunas, linhas in s.tabelas:
            ws.append(colunas)
            for c in ws[ws.max_row]:
                c.font = Font(bold=True)
                c.fill = PatternFill("solid", fgColor="E2E8F0")
            for ln in linhas:
                ws.append([v if isinstance(v, (int, float)) else _v(v) for v in ln])
    for col in "ABCDEFGHIJ":
        ws.column_dimensions[col].width = 28
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
