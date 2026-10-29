"""Tetos InvestSUS: validação com a tela Limites, semáforo, programas do MS, memória de cálculo,
alertas, importação de limites e exportação — vários municípios e exercícios."""
import io
from datetime import date, datetime, timedelta
from types import SimpleNamespace

import openpyxl
import pytest

from services import investsus_tetos as tt

_ID = iter(range(1, 10_000))


def _p(mod, sit="aprovada", acao="2E89 — Incremento Temporário PAP", fin="custeio", ind=0.0, prop=0.0,
       apr=0.0, pago=0.0, ex=2026, objeto=None, devolvido=0.0):
    return SimpleNamespace(
        id=next(_ID), numero_proposta=f"N{next(_ID)}", modalidade=mod, situacao_normalizada=sit,
        acao_orcamentaria=acao, componente=None, bloco_financiamento=None, finalidade=None, tipo_objeto=None,
        objeto=objeto, programa=None, tipo_financiamento=fin, valor_indicado=ind, valor_proposta=prop,
        valor_aprovado=apr, valor_empenhado=0.0, valor_pago=pago, valor_devolvido=devolvido, exercicio=ex)


def _reg(cat, limite, util=None, origem=tt.ORIGEM_TELA, versao=1, quando=None):
    return {"categoria": cat, "valor_limite": limite, "valor_utilizado": util,
            "saldo": None if util is None else limite - util, "origem": origem, "versao": versao,
            "registrado_em": quando or datetime(2026, 10, 29, 10), "registrado_por": "Teste"}


def _cards(props, tetos=None, registros=(), ex=2026):
    painel = tt.montar_painel(tetos or {}, props, lambda p: p.modalidade, ex, registros)
    return painel, {c["categoria"]: c for c in painel["cards"]}


# ── Rótulos da tela Limites do InvestSUS ──────────────────────────────────────

@pytest.mark.parametrize("rotulo,cat", [
    ("INCREMENTO AO CUSTEIO DE SERVIÇOS DA ATENÇÃO PRIMÁRIA À SAÚDE - COLETIVAS", "coletiva_pap"),
    ("Incremento ao custeio de serviços da Atenção Especializada à Saúde – Coletivas", "coletiva_mac"),
    ("INCREMENTO AO CUSTEIO DE SERVIÇOS DA ATENÇÃO PRIMÁRIA À SAÚDE", "individual_pap"),
    ("INCREMENTO AO CUSTEIO DE SERVIÇOS DA ATENÇÃO ESPECIALIZADA À SAÚDE", "individual_mac"),
    ("INCREMENTO AO PROGRAMA AGORA TEM ESPECIALISTAS", "programa_agora_tem_especialistas"),
    ("CUSTEIO PAP", "programa_custeio_pap"),
    ("individual_mac", "individual_mac"),
    ("RP7 MAC", "coletiva_mac"),
    ("Estruturação da rede", None),
])
def test_rotulos_da_tela_limites(rotulo, cat):
    assert tt.categoria_do_limite(rotulo) == cat


# ── Apuí/AM 2026: valores do painel conferidos (aritmética e leitura) ─────────

def test_apui_2026_saldos_do_painel_conferem_mas_ficam_pendentes_sem_tela_limites():
    tetos = {"individual_pap": {"valor": 10209109.0, "situacao": "confirmado"},
             "individual_mac": {"valor": 5675296.0, "situacao": "pendente_validacao"},
             "coletiva_pap": {"valor": 8078142.0, "situacao": "confirmado"},
             "coletiva_mac": {"valor": 5675296.0, "situacao": "pendente_validacao"}}
    mac = "2E90 — Incremento Temporário MAC"
    props = [_p("rp6", apr=2300000), _p("rp6", acao=mac, apr=1100000),
             _p("rp7", apr=5000000), _p("rp8", apr=700110),
             _p("rp7", acao=mac, apr=4000000), _p("rp8", acao=mac, apr=1086971)]
    _, c = _cards(props, tetos)
    assert {k: v["saldo"] for k, v in c.items()} == {
        "individual_pap": 7909109.0, "individual_mac": 4575296.0,
        "coletiva_pap": 2378032.0, "coletiva_mac": 588325.0}
    assert c["coletiva_mac"]["perc_utilizado"] == 89.6
    assert c["coletiva_mac"]["semaforo"]["cor"] == "amarelo"
    # Teto lido só da planilha do FNS: nada é exibido como validado.
    assert all(v["validacao"]["teto"] == tt.SITUACAO_PENDENTE for v in c.values())
    assert all(v["validacao"]["saldo"] == tt.SITUACAO_PENDENTE for v in c.values())
    # Bancada e comissão separadas, um único controle (teto não somado duas vezes).
    det = c["coletiva_pap"]["detalhe_modalidade"]
    assert det["rp7"]["comprometido"] == 5000000 and det["rp8"]["comprometido"] == 700110
    assert c["coletiva_pap"]["teto_oficial"] == 8078142.0


# ── Validação com a tela Limites do InvestSUS ─────────────────────────────────

def test_tela_limites_com_utilizado_igual_valida_teto_e_saldo():
    regs = [_reg("individual_pap", 1000.0, util=300.0)]
    painel, c = _cards([_p("rp6", apr=300)], {"individual_pap": {"valor": 1000.0}}, regs)
    card = c["individual_pap"]
    assert card["validacao"]["teto"] == "validado" and card["validacao"]["saldo"] == "validado"
    assert card["saldo"] == 700 and card["origem_teto"] == "investsus_limites"
    alerta = [a for a in tt.alertas_tetos(painel, regs) if a["tipo"] == "saldo_validado_disponivel"]
    assert alerta and "não garante indicação" in alerta[0]["mensagem"]


def test_utilizado_diferente_fica_divergente_e_gera_alerta():
    regs = [_reg("individual_pap", 1000.0, util=500.0)]
    painel, c = _cards([_p("rp6", apr=300)], {"individual_pap": {"valor": 1000.0}}, regs)
    card = c["individual_pap"]
    assert card["validacao"]["saldo"] == "divergente"
    assert card["validacao"]["divergencias"][0]["diferenca"] == 200
    assert any(a["tipo"] == "divergencia_investsus" for a in tt.alertas_tetos(painel, regs))
    assert not any(a["tipo"] == "saldo_validado_disponivel" for a in tt.alertas_tetos(painel, regs))


def test_limite_da_tela_prevalece_sobre_planilha_mac_e_registra_diferenca():
    # Ex.: MAC com acréscimo de Amazônia Legal na tela; planilha FNS traz só o valor-base.
    regs = [_reg("coletiva_mac", 7377884.8, util=0.0)]
    _, c = _cards([], {"coletiva_mac": {"valor": 5675296.0, "situacao": "pendente_validacao"}}, regs)
    card = c["coletiva_mac"]
    assert card["teto_oficial"] == 7377884.8 and card["validacao"]["teto"] == "validado"
    div = card["validacao"]["divergencias"][0]
    assert div["tipo"] == "teto_fns_x_investsus" and "valor-base" in div["mensagem"]


def test_proposta_sem_componente_impede_saldo_validado():
    regs = [_reg("individual_pap", 1000.0, util=0.0)]
    _, c = _cards([_p("rp6", acao=None, fin=None, apr=50, objeto="Custeio diverso")],
                  {"individual_pap": {"valor": 1000.0}}, regs)
    assert c["individual_pap"]["validacao"]["saldo"] == "pendente_validacao"
    assert c["individual_pap"]["a_classificar"]


def test_ultima_versao_da_tela_e_a_que_vale():
    regs = [_reg("individual_pap", 1000.0, versao=1), _reg("individual_pap", 1200.0, versao=2),
            _reg("individual_pap", 999.0, origem=tt.ORIGEM_FNS, versao=5)]
    painel, c = _cards([], {"individual_pap": {"valor": 1000.0}}, regs)
    assert c["individual_pap"]["teto_oficial"] == 1200.0
    alt = [a for a in tt.alertas_tetos(painel, regs) if a["tipo"] == "alteracao_oficial_teto"]
    assert alt and "R$ 1.200,00" in alt[0]["mensagem"]


# ── Semáforo ──────────────────────────────────────────────────────────────────

@pytest.mark.parametrize("comprometido,cor,texto", [
    (790.0, "verde", "Saldo disponível"), (800.0, "amarelo", "Próximo do esgotamento"),
    (1000.0, "vermelho", "Teto esgotado"), (1000.5, "vermelho", "Comprometido acima do teto"),
])
def test_semaforo_sempre_com_texto(comprometido, cor, texto):
    assert tt.semaforo(1000.0, comprometido) == {"cor": cor, "texto": texto}


def test_semaforo_sem_teto_e_cinza():
    assert tt.semaforo(None, 10)["cor"] == "cinza"


# ── Programas do MS separados dos tetos de emendas ────────────────────────────

def test_agora_tem_especialistas_e_custeio_pap_tem_limite_proprio():
    regs = [_reg(tt.PROGRAMA_CUSTEIO_PAP, 500.0), _reg(tt.PROGRAMA_AGORA_TEM_ESPECIALISTAS, 2000.0)]
    props = [_p("rp6", acao="Incremento ao Programa Agora Tem Especialistas", apr=1500),
             _p("programa", acao="Custeio PAP — Atenção Primária", apr=100)]
    painel, c = _cards(props, {"individual_pap": {"valor": 1000.0}, "individual_mac": {"valor": 1000.0}}, regs)
    assert c["individual_pap"]["comprometido"] == 0 and c["individual_mac"]["comprometido"] == 0
    progs = {g["categoria"]: g for g in painel["programacao"]["programas"] if g["categoria"]}
    ate = progs[tt.PROGRAMA_AGORA_TEM_ESPECIALISTAS]
    assert ate["valor_solicitado"] == 1500 and ate["saldo"] == 500
    assert progs[tt.PROGRAMA_CUSTEIO_PAP]["saldo"] == 400


def test_programas_sem_registro_aparecem_como_nao_disponivel():
    painel, _ = _cards([])
    progs = {g["categoria"]: g for g in painel["programacao"]["programas"]}
    assert progs[tt.PROGRAMA_CUSTEIO_PAP]["valor_maximo"] is None
    assert progs[tt.PROGRAMA_CUSTEIO_PAP]["semaforo"]["cor"] == "cinza"


# ── Memória de cálculo ────────────────────────────────────────────────────────

def test_memoria_justifica_inclusao_e_exclusao():
    _, c = _cards([_p("rp6", apr=300), _p("rp6", sit="cancelada", apr=200),
                   _p("rp6", acao=None, fin=None, apr=50, objeto="Custeio diverso")],
                  {"individual_pap": {"valor": 1000.0}})
    linhas = c["individual_pap"]["memoria_linhas"]
    assert [x["incluida"] for x in linhas] == [True, False, False]
    assert "uma única vez" in linhas[0]["justificativa"]
    assert linhas[1]["valor_considerado"] == 0 and "cancelada" in linhas[1]["justificativa"]
    assert "não identificado" in linhas[2]["justificativa"]


def test_devolucao_nao_recompoe_teto():
    _, c = _cards([_p("rp6", apr=300, pago=300, devolvido=100)], {"individual_pap": {"valor": 1000.0}})
    card = c["individual_pap"]
    assert card["saldo"] == 700 and card["metricas"]["valor_devolvido"] == 100
    assert any("Devoluções" in m for m in card["memoria_calculo"])


def test_exportacao_pdf_e_excel():
    painel, _ = _cards([_p("rp7", apr=300), _p("rp8", sit="rejeitada", apr=10)], {"coletiva_pap": {"valor": 1000.0}})
    ctrl = tt.controle_da_categoria(painel, "coletiva_pap")
    xlsx = tt.exportar_memoria_xlsx(ctrl, "Apuí/AM", 2026, "Teste")
    pdf = tt.exportar_memoria_pdf(ctrl, "Apuí/AM", 2026, "Teste")
    assert xlsx[:2] == b"PK" and pdf[:4] == b"%PDF"
    ws = openpyxl.load_workbook(io.BytesIO(xlsx)).active
    textos = [c for row in ws.iter_rows(values_only=True) for c in row if c]
    assert "Saldo para novas propostas" in textos and "R$ 700,00" in textos
    assert tt.controle_da_categoria(painel, tt.PROGRAMA_CUSTEIO_PAP) is not None
    assert tt.controle_da_categoria(painel, "inexistente") is None


# ── Alertas ───────────────────────────────────────────────────────────────────

def test_alertas_esgotado_aprovadas_nao_pagas_e_dados_desatualizados():
    painel, _ = _cards([_p("rp6", apr=1000)], {"individual_pap": {"valor": 1000.0}})
    agora = datetime(2026, 10, 29)
    al = tt.alertas_tetos(painel, [], ultima_sincronizacao=agora - timedelta(days=20),
                          ultima_consulta_fns=agora - timedelta(days=40), agora=agora)
    tipos = [a["tipo"] for a in al]
    assert tipos[0] == "teto_esgotado" and al[0]["nivel"] == "critico"
    assert "aprovadas_nao_pagas" in tipos
    assert tipos.count("dados_desatualizados") == 2


def test_alerta_propostas_a_regularizar():
    painel, _ = _cards([_p("rp6", acao=None, fin=None, apr=50, objeto="x")])
    assert any(a["tipo"] == "propostas_a_regularizar" for a in tt.alertas_tetos(painel, []))


# ── Importação de arquivo de limites ──────────────────────────────────────────

CSV_OK = ("exercicio;ibge;limite;valor_limite;valor_utilizado;saldo\n"
          "2026;1300144;INCREMENTO AO CUSTEIO DE SERVIÇOS DA ATENÇÃO PRIMÁRIA À SAÚDE;10.209.109,00;2.300.000,00;7.909.109,00\n"
          "2026;130014;Custeio PAP;500,00;;\n").encode("utf-8")


def test_importacao_csv_valida():
    r = tt.ler_arquivo_limites(CSV_OK, "limites.csv", "1300144", 2026)
    assert not r["erros"] and len(r["linhas"]) == 2
    assert r["linhas"][0]["categoria"] == "individual_pap" and r["linhas"][0]["valor_limite"] == 10209109
    assert r["linhas"][1]["valor_utilizado"] is None and len(r["hash_sha256"]) == 64


def test_importacao_recusa_outro_municipio_exercicio_modalidade_e_conta_errada():
    csv = ("exercicio,ibge,limite,valor_limite,valor_utilizado,saldo\n"
           "2025,1300144,Custeio PAP,10,,\n"
           "2026,1302603,Custeio PAP,10,,\n"
           "2026,1300144,Estruturação,10,,\n"
           "2026,1300144,RP6 PAP,1000,300,600\n").encode()
    r = tt.ler_arquivo_limites(csv, "x.csv", "1300144", 2026)
    assert not r["linhas"] and [e["linha"] for e in r["erros"]] == [2, 3, 4, 5]
    assert "exercício" in r["erros"][0]["erros"][0] and "IBGE" in r["erros"][1]["erros"][0]
    assert "não reconhecido" in r["erros"][2]["erros"][0] and "transcrição" in r["erros"][3]["erros"][0]


def test_importacao_xlsx_e_estrutura_invalida():
    wb = openpyxl.Workbook()
    wb.active.append(["Exercício", "IBGE", "Limite", "Valor Limite"])
    wb.active.append([2025, "1302603", "RP7 PAP", 123.45])
    buf = io.BytesIO()
    wb.save(buf)
    r = tt.ler_arquivo_limites(buf.getvalue(), "l.xlsx", "1302603", 2025)
    assert r["linhas"][0]["categoria"] == "coletiva_pap" and r["linhas"][0]["valor_limite"] == 123.45
    with pytest.raises(tt.ArquivoLimitesInvalido):
        tt.ler_arquivo_limites(b"a;b\n1;2\n", "x.csv", "1300144", 2026)
    with pytest.raises(tt.ArquivoLimitesInvalido):
        tt.ler_arquivo_limites(b"%PDF", "x.pdf", "1300144", 2026)


# ── Rotas (banco em memória): registro, importação, exercícios e municípios ───

def _usuario(mid, role="gestor"):
    from routers.auth import UserOut
    return UserOut(username="teste", nome="Teste", role=role, municipio_id=mid, municipio_ibge="1300144")


async def test_rotas_registro_importacao_e_isolamento(db_session, municipio):
    from fastapi import HTTPException
    from starlette.datastructures import UploadFile
    from models.investsus import PropostaInvestSUS, SituacaoNormalizada
    from models.municipio import Municipio
    from routers import investsus as rt

    outro = Municipio(nome="Manicoré", uf="AM", codigo_ibge="1302603")
    db_session.add(outro)
    await db_session.flush()
    db_session.add(PropostaInvestSUS(municipio_id=municipio.id, numero_proposta="P1", exercicio=2026, modalidade="rp6",
                                     acao_orcamentaria="2E89", tipo_financiamento="custeio", valor_aprovado=300.0,
                                     situacao_normalizada=SituacaoNormalizada.APROVADA))
    await db_session.commit()
    u = _usuario(municipio.id)

    with pytest.raises(HTTPException):  # sem permissão
        await rt.registrar_limite_investsus(rt.LimiteRegistro(exercicio=2026, categoria="individual_pap",
                                                              valor_limite=1000), db_session, _usuario(municipio.id, "acs"))
    with pytest.raises(HTTPException):  # conta não fecha
        await rt.registrar_limite_investsus(rt.LimiteRegistro(exercicio=2026, categoria="individual_pap",
                                                              valor_limite=1000, valor_utilizado=300, saldo=600),
                                            db_session, u)
    reg = await rt.registrar_limite_investsus(rt.LimiteRegistro(
        exercicio=2026, categoria="INCREMENTO AO CUSTEIO DE SERVIÇOS DA ATENÇÃO PRIMÁRIA À SAÚDE",
        valor_limite=1000, saldo=700, consultado_em=date(2026, 10, 29)), db_session, u)
    assert reg["categoria"] == "individual_pap" and reg["valor_utilizado"] == 300 and reg["versao"] == 1

    painel = await rt.tetos_municipais(db_session, u, 2026)
    card = {c["categoria"]: c for c in painel["cards"]}["individual_pap"]
    assert card["validacao"]["saldo"] == "validado" and card["saldo"] == 700
    assert painel["municipio"]["ibge"] == "1300144" and len(painel["historico"]) == 1

    # Prévia não grava; arquivo de outro município é recusado; confirmação grava nova versão.
    csv = "exercicio;ibge;limite;valor_limite;valor_utilizado\n2026;1300144;RP6 PAP;1100;300\n".encode()
    prev = await rt.importar_limites_investsus(db_session, u, UploadFile(io.BytesIO(csv), filename="l.csv"), 2026, False)
    assert prev["gravados"] == 0 and not prev["erros"]
    errado = "exercicio;ibge;limite;valor_limite\n2026;1302603;RP6 PAP;1\n".encode()
    prev2 = await rt.importar_limites_investsus(db_session, u, UploadFile(io.BytesIO(errado), filename="l.csv"), 2026, False)
    assert prev2["erros"] and not prev2["linhas"]
    with pytest.raises(HTTPException):
        await rt.importar_limites_investsus(db_session, u, UploadFile(io.BytesIO(errado), filename="l.csv"), 2026, True)
    ok = await rt.importar_limites_investsus(db_session, u, UploadFile(io.BytesIO(csv), filename="l.csv"), 2026, True)
    assert ok["gravados"] == 1
    painel = await rt.tetos_municipais(db_session, u, 2026)
    card = {c["categoria"]: c for c in painel["cards"]}["individual_pap"]
    assert card["teto_oficial"] == 1100 and card["limite_investsus"]["versao"] == 2
    assert card["limite_investsus"]["origem"] == "arquivo"
    assert any(a["tipo"] == "alteracao_oficial_teto" for a in painel["alertas"])

    # Outro exercício não herda o limite de 2026.
    p2025 = await rt.tetos_municipais(db_session, u, 2025)
    assert {c["categoria"]: c for c in p2025["cards"]}["individual_pap"]["teto_oficial"] is None
    assert 2026 in p2025["exercicios_disponiveis"]

    # Outro município não vê o limite nem o histórico de Apuí.
    p_outro = await rt.tetos_municipais(db_session, _usuario(outro.id), 2026)
    assert {c["categoria"]: c for c in p_outro["cards"]}["individual_pap"]["teto_oficial"] is None
    assert p_outro["historico"] == []

    # Exportação da memória.
    r = await rt.exportar_memoria_teto("individual_pap", db_session, u, 2026, "xlsx")
    assert r.body[:2] == b"PK" and "memoria_teto_individual_pap_2026.xlsx" in r.headers["content-disposition"]
    r = await rt.exportar_memoria_teto("individual_pap", db_session, u, 2026, "pdf")
    assert r.body[:4] == b"%PDF"
    with pytest.raises(HTTPException):
        await rt.exportar_memoria_teto("xpto", db_session, u, 2026, "pdf")


async def test_atualizar_fns_grava_historico_so_quando_muda(db_session, municipio, monkeypatch):
    from sqlalchemy import select
    from models.investsus import LimiteOficialInvestSUS
    from routers import investsus as rt

    valor = {"v": 1000.0}

    async def falso_fns(ibge, exercicio):
        return [{"categoria": "individual_pap", "exercicio": exercicio, "valor": valor["v"], "situacao": "confirmado",
                 "erro": None, "detalhes": {}, "documento": "Portaria", "url_planilha": "https://x/Teto_PAP.xlsx",
                 "url_pagina": "p", "data_publicacao": date(2026, 3, 3), "arquivo_modificado_em": None,
                 "hash_sha256": "h"}]
    monkeypatch.setattr(tt, "consultar_tetos", falso_fns)
    u = _usuario(municipio.id)
    await rt.atualizar_tetos_municipais(db_session, u, None)
    await rt.atualizar_tetos_municipais(db_session, u, None)  # mesmo valor: sem nova versão
    valor["v"] = 1100.0
    await rt.atualizar_tetos_municipais(db_session, u, None)
    regs = (await db_session.execute(select(LimiteOficialInvestSUS).order_by(LimiteOficialInvestSUS.versao))).scalars().all()
    assert [(r.origem, r.valor_limite, r.versao, r.arquivo_nome) for r in regs] == [
        ("planilha_fns", 1000.0, 1, "Teto_PAP.xlsx"), ("planilha_fns", 1100.0, 2, "Teto_PAP.xlsx")]
    painel = await rt.tetos_municipais(db_session, u, None)
    assert any(a["tipo"] == "alteracao_oficial_teto" for a in painel["alertas"])
    # Teto só da planilha: nunca "validado".
    card = {c["categoria"]: c for c in painel["cards"]}["individual_pap"]
    assert card["validacao"]["teto"] == "pendente_validacao"
